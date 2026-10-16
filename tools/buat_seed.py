"""Membuat supabase/02_data_awal.sql dari ekspor spreadsheet (.xlsx).

Pemakaian:  python3 tools/buat_seed.py <file.xlsx>
"""
import sys
import datetime as dt
from pathlib import Path

import openpyxl

SRC = sys.argv[1]
OUT = Path(__file__).resolve().parent.parent / "supabase" / "02_data_awal.sql"

ADMIN = [
    ("imamsyamshaq@gmail.com", "Imam"),
    ("super.arfika@gmail.com", "Abrar"),
    ("abrarakhmadr@gmail.com", "Abrar"),
]

wb = openpyxl.load_workbook(SRC, data_only=True)


def bersih(v):
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\r\n", "\n").strip()
    return s or None


def q(v):
    v = bersih(v)
    if v is None:
        return "null"
    return "'" + v.replace("'", "''") + "'"


def tabel(ws):
    head = [bersih(c.value) or "" for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(bersih(x) for x in r):
            continue
        rows.append({h: r[i] if i < len(r) else None for i, h in enumerate(head) if h})
    return rows


out = []
w = out.append
w("-- =====================================================================")
w("--  DATA AWAL (dibuat otomatis dari spreadsheet " + dt.date.today().strftime("%d/%m/%Y") + ")")
w("--  Jalankan SETELAH 01_skema.sql. Aman dijalankan ulang: data yang sama")
w("--  akan ditimpa dengan isi spreadsheet saat file ini dibuat.")
w("-- =====================================================================")
w("begin;")
w("")

# ---- admin
w("-- Admin yang boleh login dan mengubah data")
w("insert into public.admin (email, nama) values")
w(",\n".join(f"  ({q(e)}, {q(n)})" for e, n in ADMIN))
w("on conflict (email) do update set nama = excluded.nama, aktif = true;")
w("")

# ---- kategori
kat = [r for r in tabel(wb["Kategori"]) if bersih(r.get("Kode")) and bersih(r.get("Nama Sub"))]
w("-- Kategori (" + str(len(kat)) + ")")
w("insert into public.kategori (kode, nama, nama_en, label, ikon, deskripsi, deskripsi_en, target, urutan) values")
baris = []
for i, r in enumerate(kat, 1):
    target = bersih(r.get("Target"))
    baris.append("  (" + ", ".join([
        q(r["Kode"]), q(r["Nama Sub"]), q(r.get("Nama Sub (English)")), q(r.get("Label Dropdown")),
        q(r.get("Ikon")), q(r.get("Deskripsi")), q(r.get("Deskripsi (English)")),
        (str(int(float(target))) if target and target.replace(".", "", 1).isdigit() else "null"), str(i),
    ]) + ")")
w(",\n".join(baris))
w("on conflict (kode) do update set nama = excluded.nama, nama_en = excluded.nama_en, label = excluded.label,")
w("  ikon = excluded.ikon, deskripsi = excluded.deskripsi, deskripsi_en = excluded.deskripsi_en,")
w("  target = excluded.target, urutan = excluded.urutan;")
w("")

# ---- kosakata
label_ke_kode = {bersih(r.get("Label Dropdown")): bersih(r["Kode"]) for r in kat}
label_ke_kode.update({bersih(r["Nama Sub"]): bersih(r["Kode"]) for r in kat})
kos = [r for r in tabel(wb["Kosakata"]) if bersih(r.get("ID")) and bersih(r.get("Bahasa Indonesia (ID)"))]
KOL = ["id", "sub", "tipe", "indo", "en", "fusha_ar", "fusha_lat", "mesir_ar", "mesir_lat", "saudi_ar",
       "saudi_lat", "catatan", "catatan_en", "sumber", "status", "dicek_oleh"]
SHEET = ["ID", "Sub", "Tipe", "Bahasa Indonesia (ID)", "English (ENG)", "Fusha (Arab)", "Fusha (Latin)",
         "Mesir (Arab)", "Mesir (Latin)", "Saudi (Arab)", "Madinah & Makkah (Latin)", "Catatan Penggunaan",
         "Catatan (English)", "Sumber / Informan", "Status", "Dicek oleh"]
w("-- Kosakata (" + str(len(kos)) + ")")
w("insert into public.kosakata (" + ", ".join(KOL) + ") values")
baris = []
for r in kos:
    vals = []
    for k, h in zip(KOL, SHEET):
        v = r.get(h)
        if k == "sub":
            s = bersih(v) or ""
            v = label_ke_kode.get(s, s[:1].upper())
        vals.append(q(v))
    baris.append("  (" + ", ".join(vals) + ")")
w(",\n".join(baris))
w("on conflict (id) do update set " + ", ".join(f"{k} = excluded.{k}" for k in KOL[1:]) + ";")
w("")

# ---- tentang
ten = [r for r in tabel(wb["Tentang Saya"]) if bersih(r.get("Kunci"))]
kol_isi = next(h for h in ten[0] if h.startswith("Isi (") and "English" not in h)
w("-- Tentang Saya (" + str(len(ten)) + ")")
w("insert into public.tentang (kunci, bagian, panduan, isi, isi_en, urutan) values")
w(",\n".join("  (" + ", ".join([q(r["Kunci"]), q(r.get("Bagian")), q(r.get("Panduan Isi")), q(r.get(kol_isi)),
                                 q(r.get("Isi (English)")), str(i)]) + ")" for i, r in enumerate(ten, 1)))
w("on conflict (kunci) do update set bagian = excluded.bagian, panduan = excluded.panduan, isi = excluded.isi,")
w("  isi_en = excluded.isi_en, urutan = excluded.urutan;")
w("")

# ---- administrasi (hanya baris yang judulnya sudah diisi)
adm = [r for r in tabel(wb["Administrasi"]) if bersih(r.get("Judul Langkah"))]
w("-- Administrasi (" + str(len(adm)) + " langkah terisi)")
if adm:
    AK = [("id", "ID"), ("tahap", "Tahap"), ("judul", "Judul Langkah"), ("penjelasan", "Penjelasan"),
          ("dokumen", "Dokumen yang Disiapkan"), ("tempat", "Tempat / Instansi"), ("biaya", "Perkiraan Biaya"),
          ("waktu", "Perkiraan Waktu"), ("tips", "Tips dari Pengalaman"), ("sumber", "Sumber Resmi / Link"),
          ("terakhir_dicek", "Terakhir Dicek"), ("status", "Status"), ("judul_en", "Judul Langkah (ENG)"),
          ("penjelasan_en", "Penjelasan (ENG)"), ("tips_en", "Tips dari Pengalaman (ENG)")]
    w("insert into public.administrasi (" + ", ".join(k for k, _ in AK) + ", urutan) values")
    w(",\n".join("  (" + ", ".join(q(r.get(h)) for _, h in AK) + f", {i})" for i, r in enumerate(adm, 1)))
    w("on conflict (id) do update set " + ", ".join(f"{k} = excluded.{k}" for k, _ in AK[1:]) + ";")
w("")

# ---- panduan (bagian yang tampil di website)
pan = []
dalam = False
nomor = 0
for r in wb["Panduan"].iter_rows(values_only=True):
    label, isi, en = (bersih(r[1]) if len(r) > 1 else None), (bersih(r[2]) if len(r) > 2 else None), (bersih(r[3]) if len(r) > 3 else None)
    if label == "Tujuan":
        pan.append(("tujuan", "Tujuan", isi, en))
    elif label == "Kenapa 3 bahasa?":
        pan.append(("kenapa", "Kenapa 3 bahasa?", isi, en))
    elif label == "Prinsip":
        pan.append(("prinsip", "Prinsip cara baca", isi, en))
        dalam = True
    elif dalam:
        if label and isi:
            nomor += 1
            pan.append((f"baca_{nomor:02d}", label, isi, en))
        elif label and not isi:
            dalam = False
w("-- Panduan untuk website (" + str(len(pan)) + " baris)")
w("insert into public.panduan (kunci, label, isi, isi_en, urutan) values")
w(",\n".join(f"  ({q(k)}, {q(l)}, {q(i)}, {q(e)}, {n})" for n, (k, l, i, e) in enumerate(pan, 1)))
w("on conflict (kunci) do update set label = excluded.label, isi = excluded.isi, isi_en = excluded.isi_en, urutan = excluded.urutan;")
w("")

# ---- log
w("-- Catatan impor: ganti ratusan baris log 'tambah' dengan satu ringkasan")
w("delete from public.log_admin where pelaku = 'SQL editor';")
w(f"insert into public.log_admin (pelaku, aksi, tabel, baris_id, perubahan) values ('SQL editor', 'impor', 'semua', null,")
w(f"  jsonb_build_object('kategori', {len(kat)}, 'kosakata', {len(kos)}, 'tentang', {len(ten)}, 'administrasi', {len(adm)}, 'panduan', {len(pan)}));")
w("")
w("commit;")
w("")
w("-- Cek hasil (harus sama dengan angka di komentar di atas):")
w("select (select count(*) from public.kategori) as kategori, (select count(*) from public.kosakata) as kosakata,")
w("       (select count(*) from public.tentang) as tentang, (select count(*) from public.panduan) as panduan,")
w("       (select count(*) from public.admin) as admin;")

OUT.write_text("\n".join(out) + "\n", encoding="utf-8")
print("ditulis:", OUT, "| kategori", len(kat), "kosakata", len(kos), "tentang", len(ten), "adm", len(adm), "panduan", len(pan))
